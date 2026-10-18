"""Extraction per format: labels, sections, pseudo-pages and permanent failure codes. Pure unit tests (no services)."""

import io
import os
import struct
import sys
import time
import zipfile
from pathlib import Path

import docx
import pytest

import app.ingestion.extract as extract_module
from app.config import settings
from app.core.errors import ErrorCode, PermanentError
from app.ingestion.chunking import MAX_HEADER_TOKENS, MIN_TAIL_TOKENS, chunk_pages, count_tokens, embed_text
from app.ingestion.extract import (
    MAX_CHARS_PER_PAGE,
    MAX_HEADING_TOKENS,
    MAX_TOKENS_PER_PAGE,
    SECTION_TOKENS,
    extract,
)
from app.ingestion.sniff import DOCX, MAX_ENTRY_UNCOMPRESSED, MAX_UNCOMPRESSED, MD, PDF, TXT, sniff
from tests.ingestion.helpers import (
    PAGE_TEXTS,
    encrypt_pdf,
    label_pdf_pages,
    make_docx,
    make_image_only_pdf,
    make_mixed_pdf,
    make_pdf,
)


def _file(tmp_path: Path, name: str, data: bytes | str) -> Path:
    path = tmp_path / name
    path.write_bytes(data.encode() if isinstance(data, str) else data)
    return path


def _paragraphs(n: int, words: int = 60) -> str:
    return "\n\n".join(f"Paragraph {i}. " + " ".join(f"word{i}x{j}" for j in range(words)) for i in range(n))


# --- PDF ----------------------------------------------------------------------------------------------------------


def test_pdf_one_page_per_physical_page_with_default_labels(tmp_path: Path) -> None:
    doc = extract(_file(tmp_path, "a.pdf", make_pdf()), PDF)
    assert [p.page_no for p in doc.pages] == [1, 2, 3]
    assert [p.label for p in doc.pages] == ["p. 1", "p. 2", "p. 3"]
    assert "LRK-4471" in doc.pages[1].text
    assert all(p.heading_path is None for p in doc.pages)
    assert doc.warnings == []


def test_pdf_uses_the_documents_own_page_labels(tmp_path: Path) -> None:
    data = label_pdf_pages(make_pdf([*PAGE_TEXTS, "Appendix text that is long enough to count."]))
    doc = extract(_file(tmp_path, "a.pdf", data), PDF)
    assert [p.label for p in doc.pages] == ["p. i", "p. ii", "p. 1", "p. 2"]
    assert [p.page_no for p in doc.pages] == [1, 2, 3, 4]   # page_no stays physical for the page viewer


def test_password_protected_pdf_is_encrypted_pdf(tmp_path: Path) -> None:
    data = encrypt_pdf(make_pdf(), user_password="s3cret")
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, "a.pdf", data), PDF)
    assert err.value.code == ErrorCode.ENCRYPTED_PDF


def test_pdf_encrypted_with_empty_user_password_is_readable(tmp_path: Path) -> None:
    data = encrypt_pdf(make_pdf(), user_password="")   # "permissions only" encryption
    doc = extract(_file(tmp_path, "a.pdf", data), PDF)
    assert "18 PTO days" in doc.pages[0].text


def test_image_only_pdf_is_no_extractable_text(tmp_path: Path) -> None:
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, "scan.pdf", make_image_only_pdf()), PDF)
    assert err.value.code == ErrorCode.NO_EXTRACTABLE_TEXT
    assert "OCR is not supported" in err.value.message


def test_partly_scanned_pdf_is_indexed_with_a_warning(tmp_path: Path) -> None:
    doc = extract(_file(tmp_path, "mixed.pdf", make_mixed_pdf()), PDF)
    assert len(doc.pages) == 2                      # the scanned page is kept so page numbers stay true
    assert doc.warnings == ["pages_without_text:1"]


def test_pdf_whose_only_text_is_short_is_not_mistaken_for_a_scan(tmp_path: Path) -> None:
    doc = extract(_file(tmp_path, "receipt.pdf", make_pdf(["Invoice #1234 paid"])), PDF)
    assert doc.pages[0].text.strip() == "Invoice #1234 paid"
    assert doc.warnings == []                       # its only page has text: nothing chat can't see


def test_short_text_pdf_warns_only_about_its_blank_pages(tmp_path: Path) -> None:
    doc = extract(_file(tmp_path, "receipt.pdf", make_pdf(["Invoice #1234 paid", ""])), PDF)
    assert doc.warnings == ["pages_without_text:1"]


def test_near_empty_page_among_full_pages_still_counts_as_without_text(tmp_path: Path) -> None:
    doc = extract(_file(tmp_path, "a.pdf", make_pdf([PAGE_TEXTS[0], "Figure 3"])), PDF)
    assert doc.warnings == ["pages_without_text:1"]   # a stray label on a scanned page is not a text layer


@pytest.mark.parametrize("stamps", [["1", "2", "3"], ["ABC000101", "ABC000102", "ABC000103"]])
def test_scan_whose_only_text_is_stamped_page_or_bates_numbers_is_no_extractable_text(
    tmp_path: Path, stamps: list[str],
) -> None:
    """Scanned productions are often stamped with page or Bates numbers: a text layer of a few characters on every
    page and nothing more. That is still a scan, so it fails fast with the OCR hint instead of indexing the stamps."""
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, "scan.pdf", make_pdf(stamps)), PDF)
    assert err.value.code == ErrorCode.NO_EXTRACTABLE_TEXT


def test_too_many_pages(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "MAX_PAGES", 2)
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, "a.pdf", make_pdf()), PDF)
    assert err.value.code == ErrorCode.TOO_MANY_PAGES


def test_corrupt_pdf(tmp_path: Path) -> None:
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, "a.pdf", b"%PDF-1.7\nthis is not really a pdf"), PDF)
    assert err.value.code == ErrorCode.CORRUPT_FILE


# --- DOCX ---------------------------------------------------------------------------------------------------------


def test_docx_sections_follow_headings_with_breadcrumbs(tmp_path: Path) -> None:
    doc = extract(_file(tmp_path, "a.docx", make_docx()), DOCX)
    labels = [p.label for p in doc.pages]
    # "Leave Policy" has no text of its own (a sub-heading follows), so it only appears in its children's path.
    assert labels == ["Intro", "§ Paid Time Off", "§ Parental Leave", "§ Security"]
    by_label = {p.label: p for p in doc.pages}
    assert by_label["Intro"].heading_path is None
    assert by_label["§ Paid Time Off"].heading_path == "Leave Policy > Paid Time Off"
    assert by_label["§ Parental Leave"].heading_path == "Leave Policy > Parental Leave"   # H2 replaces H2
    assert by_label["§ Security"].heading_path == "Security"                              # H1 resets the path
    assert [p.page_no for p in doc.pages] == [1, 2, 3, 4]


def test_docx_tables_are_read_in_document_order_as_pipe_rows(tmp_path: Path) -> None:
    pto = extract(_file(tmp_path, "a.docx", make_docx()), DOCX).pages[1].text
    assert "Tenure | Days | Carry-over" in pto
    assert "3+ years | 22 | 5" in pto
    assert pto.index("18 PTO days") < pto.index("Tenure | Days") < pto.index("Unused days expire")


def test_long_docx_section_is_split_into_parts(tmp_path: Path) -> None:
    doc = extract(_file(tmp_path, "a.docx", make_docx(long_section_words=1500)), DOCX)
    security = [p for p in doc.pages if p.label.startswith("§ Security")]
    assert len(security) >= 2
    assert [p.label for p in security][:2] == ["§ Security (part 1)", "§ Security (part 2)"]
    assert all(p.heading_path == "Security" for p in security)
    assert all(p.text.startswith("Security\n\n") for p in security)       # every part carries its heading
    assert all(50 <= count_tokens(p.text) <= SECTION_TOKENS + MIN_TAIL_TOKENS for p in security)


def test_long_heading_is_truncated_in_the_label(tmp_path: Path) -> None:
    heading = "A very long heading " * 10
    doc = extract(_file(tmp_path, "a.md", f"# {heading}\n\nBody text."), MD)
    assert len(doc.pages[0].label) <= len("§ ") + 60
    assert doc.pages[0].label.endswith("…")


def test_corrupt_docx(tmp_path: Path) -> None:
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, "a.docx", b"PK\x03\x04 broken zip"), DOCX)
    assert err.value.code == ErrorCode.CORRUPT_FILE


W_XMLNS = b'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _docx_replacing(parts: dict[str, bytes]) -> bytes:
    """make_docx() with some parts' XML swapped, plus stored random padding that keeps the archive under the upload's
    compression-ratio limit: how a crafted upload would be built."""
    source = zipfile.ZipFile(io.BytesIO(make_docx()))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for info in source.infolist():
            zf.writestr(info.filename, parts.get(info.filename) or source.read(info.filename))
        zf.writestr("word/media/padding.bin", os.urandom(512 * 1024), compress_type=zipfile.ZIP_STORED)
    return buf.getvalue()


def _fail_if_parsed(*_args: object, **_kwargs: object) -> None:
    raise AssertionError("python-docx opened an archive that should have been refused first")


@pytest.mark.parametrize(("part", "root", "element"), [
    ("word/document.xml", (b"<w:document %s><w:body>" % W_XMLNS, b"</w:body></w:document>"), b"<w:p/>"),
    ("word/styles.xml", (b"<w:styles %s>" % W_XMLNS, b"</w:styles>"), b"<w:style/>"),
])
def test_docx_of_millions_of_tiny_elements_is_refused_before_python_docx_parses_it(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, part: str, root: tuple[bytes, bytes], element: bytes,
) -> None:
    """2M empty elements are 12-20 MB of XML, under every byte limit the upload checks, yet python-docx would turn
    them into an lxml tree of several hundred MB. Any part counts, not just the body: python-docx parses the styles,
    numbering, headers and footers as soon as it opens the file."""
    data = _docx_replacing({part: root[0] + element * 2_000_000 + root[1]})
    assert sniff(io.BytesIO(data), "bomb.docx").mime_type == DOCX          # passes every upload check
    monkeypatch.setattr(docx, "Document", _fail_if_parsed)
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, "bomb.docx", data), DOCX)
    assert err.value.code == ErrorCode.UNSAFE_ARCHIVE


NAMESPACE_DENSE = b"<w:p %s/>" % b" ".join(b'xmlns:n%d="u"' % i for i in range(400))


@pytest.mark.parametrize(("unit", "count"), [
    (b"<!---->", 2_000_000),
    (b"<?a?>", 2_000_000),
    (NAMESPACE_DENSE, 2_000),       # 400 declarations on each element
], ids=["comments", "processing-instructions", "namespace-declarations"])
def test_docx_of_millions_of_comments_pis_or_namespace_declarations_is_refused_too(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, unit: bytes, count: int,
) -> None:
    """python-docx's parser keeps comments, processing instructions and namespace declarations in its tree, each
    costing about what an element does: 29 MB of '<?a?>' in a 591 KB upload cost the worker +809 MB."""
    body = b"<w:p><w:r><w:t>Employees receive 18 PTO days per year.</w:t></w:r></w:p>" + unit * count
    document = b"<w:document %s><w:body>%s</w:body></w:document>" % (W_XMLNS, body)
    data = _docx_replacing({"word/document.xml": document})
    assert sniff(io.BytesIO(data), "bomb.docx").mime_type == DOCX          # passes every upload check
    monkeypatch.setattr(docx, "Document", _fail_if_parsed)
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, "bomb.docx", data), DOCX)
    assert err.value.code == ErrorCode.UNSAFE_ARCHIVE


def _docx_with_unreadable_extra_entry(defect: str) -> bytes:
    """make_docx() plus an entry that nothing in the package refers to and that zipfile cannot read."""
    junk = b"JUNKDATA" * 64
    buf = io.BytesIO(make_docx())
    with zipfile.ZipFile(buf, "a") as zf:
        zf.writestr("zzjunk.bin", junk, compress_type=zipfile.ZIP_STORED)
    data = bytearray(buf.getvalue())
    if defect == "bad CRC":
        data[data.index(junk)] ^= 0xFF
    else:  # compression method 9 (Deflate64), which zipfile does not support: set in the local and central headers
        local = zipfile.ZipFile(io.BytesIO(bytes(data))).getinfo("zzjunk.bin").header_offset
        central = data.rindex(b"zzjunk.bin") - 46
        for offset in (local + 8, central + 10):
            data[offset:offset + 2] = struct.pack("<H", 9)
    return bytes(data)


@pytest.mark.parametrize("defect", ["bad CRC", "unsupported compression"])
def test_unreadable_entry_python_docx_never_opens_does_not_sink_the_document(tmp_path: Path, defect: str) -> None:
    """The XML count reads every entry, but python-docx opens only the parts the package refers to. An entry outside
    those that zipfile cannot read must not fail a document python-docx extracts fine."""
    data = _docx_with_unreadable_extra_entry(defect)
    assert sniff(io.BytesIO(data), "a.docx").mime_type == DOCX
    doc = extract(_file(tmp_path, "a.docx", data), DOCX)
    clean = extract(_file(tmp_path, "clean.docx", make_docx()), DOCX)
    assert [p.label for p in doc.pages] == [p.label for p in clean.pages]


def test_docx_bomb_in_utf32_is_counted_too(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """python-docx's parser reads UTF-16 and UTF-32 XML too. A count that decoded the part differently would see no
    elements in a UTF-32 body at all, while python-docx built 1M of them."""
    body = "<w:p/>" * 1_000_000
    xml = f"<w:document {W_XMLNS.decode()}><w:body>{body}</w:body></w:document>".encode("utf-32")
    data = _docx_replacing({"word/document.xml": xml})
    assert sniff(io.BytesIO(data), "bomb.docx").mime_type == DOCX          # 24 MB: under the per-entry byte cap
    monkeypatch.setattr(docx, "Document", _fail_if_parsed)
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, "bomb.docx", data), DOCX)
    assert err.value.code == ErrorCode.UNSAFE_ARCHIVE


def _docx_with_zero_entries(sizes: list[int]) -> bytes:
    """make_docx() plus deflated zero-filled entries: tiny on disk, but zipfile hands back every byte on read."""
    buf = io.BytesIO(make_docx())
    with zipfile.ZipFile(buf, "a", zipfile.ZIP_DEFLATED) as zf:
        for i, size in enumerate(sizes):
            zf.writestr(f"word/media/zeros{i}.bin", bytes(size))
    return buf.getvalue()


@pytest.mark.parametrize("sizes", [
    [MAX_ENTRY_UNCOMPRESSED + 1],                                   # one entry over the per-entry cap
    [MAX_UNCOMPRESSED // 3 + 1] * 3,                                # each under it, together over the total
], ids=["entry", "total"])
def test_docx_over_the_upload_caps_is_refused_before_any_entry_is_read(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, sizes: list[int],
) -> None:
    """The upload's sniff refuses such an archive, but the worker extracts stored objects without sniffing them again,
    and a DOCX stored under the old 200 MB total cap can still be reprocessed. The XML count read each entry whole
    (a 90 MB entry peaked at ~200 MB) before checking anything: the entry sizes in the zip headers come first now."""
    data = _docx_with_zero_entries(sizes)
    read: list[str] = []
    real_read = zipfile.ZipFile.read

    def recording_read(self: zipfile.ZipFile, name: str | zipfile.ZipInfo, pwd: bytes | None = None) -> bytes:
        read.append(name.filename if isinstance(name, zipfile.ZipInfo) else name)
        return real_read(self, name, pwd)

    monkeypatch.setattr(zipfile.ZipFile, "read", recording_read)
    monkeypatch.setattr(docx, "Document", _fail_if_parsed)
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, "big.docx", data), DOCX)
    assert err.value.code == ErrorCode.UNSAFE_ARCHIVE
    assert not [name for name in read if name.startswith("word/media/zeros")]


def _docx_with_vertical_merge_chain(rows: int) -> bytes:
    """A one-column table whose first cell starts a vertical merge that every later row continues."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = docx.Document()
    d.add_paragraph("Employees receive 18 PTO days per year.")
    for i, row in enumerate(d.add_table(rows=rows, cols=1).rows):
        merge = OxmlElement("w:vMerge")
        if i == 0:
            merge.set(qn("w:val"), "restart")
        row.cells[0]._tc.get_or_add_tcPr().append(merge)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_that_overflows_the_python_stack_is_corrupt_not_retryable(tmp_path: Path) -> None:
    """python-docx resolves a continued vertical merge by recursing up the rows. A 1,100-row chain (37 KB) burned
    ~27 s of worker CPU and then escaped as an unexpected RecursionError: 'Reprocess to try again', forever. The
    stack is lowered here so a short chain shows the same thing quickly."""
    path = _file(tmp_path, "merged.docx", _docx_with_vertical_merge_chain(300))
    depth, frame = 0, sys._getframe()
    while frame is not None:
        depth, frame = depth + 1, frame.f_back
    limit = sys.getrecursionlimit()
    sys.setrecursionlimit(depth + 120)
    try:
        with pytest.raises(PermanentError) as err:
            extract(path, DOCX)
    finally:
        sys.setrecursionlimit(limit)
    assert err.value.code == ErrorCode.CORRUPT_FILE


def test_docx_part_with_a_doctype_is_refused(tmp_path: Path) -> None:
    """Word never writes a DOCTYPE. One in a .docx can only be there to declare entities, which the counting parser
    cannot see into, so the archive is refused rather than counted wrong."""
    document_xml = zipfile.ZipFile(io.BytesIO(make_docx())).read("word/document.xml")
    with_doctype = document_xml.replace(b"?>", b'?><!DOCTYPE w:document [<!ENTITY a "x">]>', 1)
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, "a.docx", _docx_replacing({"word/document.xml": with_doctype})), DOCX)
    assert err.value.code == ErrorCode.UNSAFE_ARCHIVE


def test_docx_style_lookups_do_not_grow_with_paragraphs_times_styles(tmp_path: Path) -> None:
    """python-docx's paragraph.style scans every style in styles.xml, in Python, on every call. A styles.xml padded
    with thousands of styles (well under the XML node limit) then made a short document cost minutes of worker time;
    here 300 paragraphs x 20k styles took ~13 s."""
    styles = (b"<w:styles %s>" % W_XMLNS
              + b'<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>'
              + b'<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/></w:style>'
              + b'<w:style w:type="character"/>' * 20_000 + b"</w:styles>")
    heading = b'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Leave Policy</w:t></w:r></w:p>'
    body = b"<w:p><w:r><w:t>Employees receive 18 PTO days per year.</w:t></w:r></w:p>" * 300
    document = b"<w:document %s><w:body>%s%s</w:body></w:document>" % (W_XMLNS, heading, body)
    path = _file(tmp_path, "a.docx", _docx_replacing({"word/styles.xml": styles, "word/document.xml": document}))
    started = time.perf_counter()
    doc = extract(path, DOCX)
    assert time.perf_counter() - started < 5
    assert {p.heading_path for p in doc.pages} == {"Leave Policy"}   # 'heading 1' is still read as Heading 1


# --- TXT / MD -------------------------------------------------------------------------------------------------------


def test_short_txt_is_one_pseudo_page(tmp_path: Path) -> None:
    doc = extract(_file(tmp_path, "a.txt", "Employees receive 18 PTO days per year."), TXT)
    assert [(p.page_no, p.label) for p in doc.pages] == [(1, "part 1")]


def test_long_txt_becomes_pseudo_pages_cut_at_paragraphs(tmp_path: Path) -> None:
    text = _paragraphs(40)
    doc = extract(_file(tmp_path, "a.txt", text), TXT)
    assert len(doc.pages) >= 2
    assert [p.label for p in doc.pages] == [f"part {n}" for n in range(1, len(doc.pages) + 1)]
    assert all(count_tokens(p.text) <= SECTION_TOKENS + MIN_TAIL_TOKENS for p in doc.pages)
    assert all(p.text.startswith("Paragraph ") for p in doc.pages)   # every page starts at a paragraph boundary


def test_markdown_headings_make_sections_but_code_comments_do_not(tmp_path: Path) -> None:
    md = (
        "Preface line.\n\n"
        "# Setup\n\nInstall it.\n\n```bash\n# not a heading\npip install vault\n```\n\n"
        "## Config ##\n\nSet VAULT_KEY.\n"
    )
    doc = extract(_file(tmp_path, "a.md", md), MD)
    assert [p.label for p in doc.pages] == ["Intro", "§ Setup", "§ Config"]
    assert "# not a heading" in doc.pages[1].text
    assert doc.pages[2].heading_path == "Setup > Config"


def test_markdown_without_headings_is_treated_as_plain_text(tmp_path: Path) -> None:
    doc = extract(_file(tmp_path, "a.md", "Just a note.\n\nNo headings here."), MD)
    assert [p.label for p in doc.pages] == ["part 1"]


def test_nul_characters_are_stripped(tmp_path: Path) -> None:
    doc = extract(_file(tmp_path, "a.txt", b"Fee is \x00$60k.\x00"), TXT)
    assert doc.pages[0].text == "Fee is $60k."


def test_whitespace_only_text_is_no_extractable_text(tmp_path: Path) -> None:
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, "a.txt", "  \n\n \t "), TXT)
    assert err.value.code == ErrorCode.NO_EXTRACTABLE_TEXT


def test_special_token_text_does_not_crash_the_tokenizer(tmp_path: Path) -> None:
    doc = extract(_file(tmp_path, "a.txt", "Prompt marker <|endoftext|> appears in this log."), TXT)
    assert "<|endoftext|>" in doc.pages[0].text


@pytest.mark.parametrize(("name", "mime"), [("a.txt", TXT), ("a.md", MD), ("a.docx", DOCX)])
def test_non_pdf_with_more_text_than_max_pages_worth_is_too_much_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, name: str, mime: str,
) -> None:
    """Only PDFs have a page count to cap; a 25 MB TXT/MD/DOCX must not become ~15k chunks."""
    monkeypatch.setattr(settings, "MAX_PAGES", 1)   # the text limit scales with the page limit
    big = make_docx(long_section_words=2000) if mime == DOCX else "# Notes\n\n" + _paragraphs(100)
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, name, big), mime)
    assert err.value.code == ErrorCode.TOO_MUCH_TEXT
    assert "characters" in err.value.message
    small = make_docx() if mime == DOCX else "# Notes\n\nEmployees receive 18 PTO days per year."
    assert extract(_file(tmp_path, name, small), mime).pages          # well under the same limit


def _docx_of(text: str) -> bytes:
    d = docx.Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize(("name", "mime"), [("a.txt", TXT), ("a.md", MD), ("a.docx", DOCX)])
def test_token_dense_text_under_the_character_limit_is_still_too_much_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, name: str, mime: str,
) -> None:
    """What the limit protects (chunks, embedding calls, the worker's time limit) grows with tokens, not characters.
    English runs ~4 characters per token, but Chinese 1.3-2.4 tokens per character: 2.5M characters of it passed the
    character limit as up to ~6M tokens and ~20k chunks."""
    monkeypatch.setattr(settings, "MAX_PAGES", 1)
    sentence = "雇员每年享有十八天带薪休假，未使用的假期在年底作废。"   # "18 paid days off a year; unused ones expire"
    text = (sentence * 200)[: MAX_CHARS_PER_PAGE - 1]
    assert count_tokens(text) > MAX_TOKENS_PER_PAGE                         # under one limit, over the other
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, name, _docx_of(text) if mime == DOCX else text), mime)
    assert err.value.code == ErrorCode.TOO_MUCH_TEXT
    assert "tokens" in err.value.message


def test_token_limit_is_counted_in_bounded_slices_and_stops_once_over(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """tiktoken builds the whole list of tokens before it can be counted. In one call, 2.5M emoji (under the character
    limit) held ~5M tokens, +500 MB, just to refuse the file: the refusal itself was the memory spike."""
    seen: list[int] = []

    def recording_count_tokens(text: str) -> int:
        seen.append(len(text))
        return count_tokens(text)

    monkeypatch.setattr(extract_module, "count_tokens", recording_count_tokens)
    text = "😀" * (settings.MAX_PAGES * MAX_CHARS_PER_PAGE - 1)    # ~2 tokens each: ~8x the token limit
    with pytest.raises(PermanentError) as err:
        extract(_file(tmp_path, "a.txt", text), TXT)
    assert err.value.code == ErrorCode.TOO_MUCH_TEXT
    assert max(seen) <= 100_000                     # never the whole text in one call
    assert sum(seen) < len(text) // 2               # and no further than needed to pass the limit


# --- Over-long "headings" ---------------------------------------------------------------------------------------------


def _docx_with_heading_styled_paragraph(words: int) -> bytes:
    """A real-world Word quirk: a whole paragraph of body text formatted with the 'Heading 1' style."""
    import io

    import docx

    d = docx.Document()
    d.add_heading("Terms", level=1)
    d.add_paragraph("Short intro.")
    d.add_heading(" ".join(f"Clause{i} applies to every supplier." for i in range(words)), level=1)
    d.add_paragraph("Closing paragraph.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_paragraph_styled_as_heading_is_body_text_so_embedding_inputs_stay_small(tmp_path: Path) -> None:
    doc = extract(_file(tmp_path, "a.docx", _docx_with_heading_styled_paragraph(1500)), DOCX)
    assert all(p.heading_path in (None, "Terms") for p in doc.pages)             # not a breadcrumb
    assert "Clause1499 applies" in "\n".join(p.text for p in doc.pages)           # ...but its text is kept
    drafts = chunk_pages(doc.pages, title="contract", model="m", dims=8)
    longest = max(count_tokens(embed_text(d.header, d.content)) for d in drafts)
    assert longest <= settings.CHUNK_TOKENS + MIN_TAIL_TOKENS + MAX_HEADER_TOKENS   # far below the 8,192 limit


def test_markdown_heading_line_over_the_limit_is_body_text(tmp_path: Path) -> None:
    long_line = " ".join(f"word{i}" for i in range(MAX_HEADING_TOKENS * 2))
    doc = extract(_file(tmp_path, "a.md", f"# Setup\n\nInstall it.\n\n## {long_line}\n\nMore text.\n"), MD)
    assert [p.heading_path for p in doc.pages] == ["Setup"]
    assert long_line in doc.pages[0].text
