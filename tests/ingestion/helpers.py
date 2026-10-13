"""Fixture builders (documents generated in-test, nothing binary checked in) and small sync-DB helpers."""

import io
import uuid
from typing import Any

from fpdf import FPDF
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from pypdf.constants import PageLabelStyle
from sqlalchemy import insert, select, text

from app.core.db import sync_session
from app.models import Chunk, Document, ProcessingRun
from app.services import state

PAGE_TEXTS = [
    "Larkspur Analytics employee handbook. Employees receive 18 PTO days per year after onboarding.",
    "Security incident LRK-4471 was resolved in March. The root cause was an expired certificate.",
    "The Brightline MSA notice period for termination for convenience is 60 days, per clause 9.2.",
]


# --- Documents ---------------------------------------------------------------------------------------------------


def make_pdf(pages: list[str] = PAGE_TEXTS) -> bytes:
    pdf = FPDF()
    pdf.set_font("helvetica", size=12)
    for body in pages:
        pdf.add_page()
        if body:
            pdf.multi_cell(0, 6, body)
    return bytes(pdf.output())


def make_image_only_pdf(pages: int = 2) -> bytes:
    """What a scanner produces: each page is a picture of text, with no text layer."""
    picture = Image.new("RGB", (400, 200), "white")
    ImageDraw.Draw(picture).rectangle((20, 20, 380, 60), fill="black")
    pdf = FPDF()
    for _ in range(pages):
        pdf.add_page()
        pdf.image(picture, x=10, y=10, w=150)
    return bytes(pdf.output())


def make_mixed_pdf() -> bytes:
    """Page 1 has text, page 2 is image-only (a partly scanned document)."""
    picture = Image.new("RGB", (200, 100), "white")
    pdf = FPDF()
    pdf.set_font("helvetica", size=12)
    pdf.add_page()
    pdf.multi_cell(0, 6, PAGE_TEXTS[0])
    pdf.add_page()
    pdf.image(picture, x=10, y=10, w=100)
    return bytes(pdf.output())


def _rewrite(data: bytes) -> PdfWriter:
    return PdfWriter(clone_from=PdfReader(io.BytesIO(data)))


def _to_bytes(writer: PdfWriter) -> bytes:
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def encrypt_pdf(data: bytes, user_password: str) -> bytes:
    writer = _rewrite(data)
    writer.encrypt(user_password, owner_password="owner-secret", algorithm="RC4-128")  # RC4: no crypto lib needed
    return _to_bytes(writer)


def label_pdf_pages(data: bytes) -> bytes:
    """Front matter numbered i, ii; body numbered from 1 — like most books and reports."""
    writer = _rewrite(data)
    writer.set_page_label(0, 1, PageLabelStyle.LOWERCASE_ROMAN)
    writer.set_page_label(2, len(writer.pages) - 1, PageLabelStyle.DECIMAL, start=1)
    return _to_bytes(writer)


def make_docx(long_section_words: int = 0) -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("Confidential draft prepared for the leadership team.")
    d.add_heading("Leave Policy", level=1)
    d.add_heading("Paid Time Off", level=2)
    d.add_paragraph("Employees receive 18 PTO days per year.")
    table = d.add_table(rows=2, cols=3)
    for r, row in enumerate([("Tenure", "Days", "Carry-over"), ("3+ years", "22", "5")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    d.add_paragraph("Unused days expire at year end.")
    d.add_heading("Parental Leave", level=2)
    d.add_paragraph("Sixteen weeks fully paid.")
    d.add_heading("Security", level=1)
    d.add_paragraph("Incident LRK-4471 was resolved. " + "Details follow. " * long_section_words)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# --- Database ----------------------------------------------------------------------------------------------------


def get_doc(document_id: str | uuid.UUID) -> Document:
    with sync_session() as s:
        return s.get(Document, uuid.UUID(str(document_id)))


def chunks_of(document_id: str | uuid.UUID) -> list[Chunk]:
    with sync_session() as s:
        return list(s.scalars(select(Chunk).where(Chunk.document_id == uuid.UUID(str(document_id)))
                              .order_by(Chunk.chunk_index)))


def runs_of(document_id: str | uuid.UUID, task: str = "ingest") -> list[ProcessingRun]:
    with sync_session() as s:
        return list(s.scalars(select(ProcessingRun)
                              .where(ProcessingRun.document_id == uuid.UUID(str(document_id)),
                                     ProcessingRun.task == task)
                              .order_by(ProcessingRun.id)))


def requeue(document_id: str | uuid.UUID, scope: str = "all") -> uuid.UUID:
    """What POST /reprocess does to the row (new run_id), without publishing a task."""
    run_id = uuid.uuid4()
    with sync_session() as s:
        s.execute(state.requeue_stmt(uuid.UUID(str(document_id)), run_id, scope=scope))
    return run_id


def insert_running_attempts(document_id: str | uuid.UUID, run_id: uuid.UUID, n: int, task: str = "ingest") -> None:
    """Simulate attempts whose worker was killed mid-task: rows left in 'running'."""
    with sync_session() as s:
        s.execute(insert(ProcessingRun), [
            {"run_id": run_id, "document_id": uuid.UUID(str(document_id)), "task": task, "attempt": i + 1,
             "status": "running"}
            for i in range(n)
        ])


def count(sql: str, **params: Any) -> int:
    with sync_session() as s:
        return int(s.scalar(text(sql), params) or 0)
